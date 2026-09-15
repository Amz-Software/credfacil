"""Geração de planilhas (.xlsx) a partir dos dados de vendas."""

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from vendas.utils import formatar_cpf, formatar_telefone

CONTENT_TYPE_XLSX = (
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
)

COLUNAS_CLIENTES = [
    ('Nome completo', 40),
    ('CPF', 20),
    ('Telefone', 20),
    ('Cidade da solicitação', 30),
]

COR_CABECALHO = '366092'
LINHA_CABECALHO = 1


def _aplicar_cabecalho(ws) -> None:
    """Escreve e estiliza a linha de cabeçalho da planilha de clientes."""
    fonte = Font(bold=True, color='FFFFFF')
    preenchimento = PatternFill(
        start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type='solid'
    )
    borda = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )
    alinhamento = Alignment(horizontal='center', vertical='center')

    for coluna, (titulo, largura) in enumerate(COLUNAS_CLIENTES, start=1):
        celula = ws.cell(row=LINHA_CABECALHO, column=coluna, value=titulo)
        celula.font = fonte
        celula.fill = preenchimento
        celula.border = borda
        celula.alignment = alinhamento
        ws.column_dimensions[get_column_letter(coluna)].width = largura


def gerar_planilha_clientes(clientes) -> HttpResponse:
    """Monta o .xlsx dos clientes recebidos e devolve o download.

    `clientes` é qualquer iterável de `Cliente` — normalmente o queryset já
    filtrado pela listagem de solicitações.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = 'Clientes'

    _aplicar_cabecalho(ws)

    for linha, cliente in enumerate(clientes, start=LINHA_CABECALHO + 1):
        ws.cell(row=linha, column=1, value=cliente.nome)
        ws.cell(row=linha, column=2, value=formatar_cpf(cliente.cpf))
        ws.cell(row=linha, column=3, value=formatar_telefone(cliente.telefone))
        ws.cell(row=linha, column=4, value=cliente.cidade)

    ws.freeze_panes = ws.cell(row=LINHA_CABECALHO + 1, column=1)

    resposta = HttpResponse(content_type=CONTENT_TYPE_XLSX)
    nome_arquivo = f'clientes_{timezone.localtime():%Y%m%d_%H%M}.xlsx'
    resposta['Content-Disposition'] = f'attachment; filename="{nome_arquivo}"'
    wb.save(resposta)
    return resposta
