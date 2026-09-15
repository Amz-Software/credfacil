"""Exportação em Excel dos clientes e filtro múltiplo de loja na listagem."""

import io
from datetime import date

from django.contrib.auth.models import Permission
from django.contrib.sessions.backends.db import SessionStore
from django.contrib.messages.storage.fallback import FallbackStorage
from django.test import RequestFactory, TestCase
from openpyxl import load_workbook

from accounts.models import User
from vendas.filtros import filtrar_clientes_solicitacao, lojas_selecionadas
from vendas.models import Cliente, ComprovantesCliente, Loja
from vendas.utils import formatar_cpf, formatar_telefone
from vendas.views import ClienteExportExcelView, ClienteListView


def _cliente(nome, loja, cidade='Belém', cpf='12345678900', telefone='91999998888'):
    return Cliente.objects.create(
        nome=nome, telefone=telefone, cpf=cpf, loja=loja,
        nascimento=date(1990, 1, 1), rg='123', cep='66000000',
        endereco='Rua A', bairro='Centro', cidade=cidade,
        comprovantes=ComprovantesCliente.objects.create(),
    )


class FormatadoresTests(TestCase):
    def test_formata_cpf_de_11_digitos(self):
        self.assertEqual(formatar_cpf('12345678900'), '123.456.789-00')

    def test_mantem_cpf_de_tamanho_inesperado(self):
        self.assertEqual(formatar_cpf('123'), '123')

    def test_formata_celular_de_11_digitos(self):
        self.assertEqual(formatar_telefone('91988887777'), '(91) 98888-7777')

    def test_formata_fixo_de_10_digitos(self):
        self.assertEqual(formatar_telefone('9132221111'), '(91) 3222-1111')

    def test_telefone_vazio_vira_string_vazia(self):
        self.assertEqual(formatar_telefone(None), '')


class FiltroLojaMultiploTests(TestCase):
    def setUp(self):
        self.factory = RequestFactory()
        self.loja_a = Loja.objects.create(nome='Loja A')
        self.loja_b = Loja.objects.create(nome='Loja B')
        self.loja_c = Loja.objects.create(nome='Loja C')

        self.user = User.objects.create_user(
            username='analista', email='analista@teste.com', password='x', loja=self.loja_a,
        )
        self.user.user_permissions.add(
            Permission.objects.get(codename='view_all_analise_credito'),
            Permission.objects.get(codename='view_cliente'),
        )
        self.user = User.objects.get(pk=self.user.pk)  # limpa o cache de permissões

        self.cliente_a = _cliente('Ana Silva', self.loja_a)
        self.cliente_b = _cliente('Bruno Costa', self.loja_b)
        self.cliente_c = _cliente('Carla Dias', self.loja_c)

    def _request(self, params):
        request = self.factory.get('/clientes/', params)
        request.user = self.user
        request.session = SessionStore()
        request.session['loja_id'] = self.loja_a.id
        return request

    def test_lojas_selecionadas_ignora_valores_vazios(self):
        request = self._request({'loja': ['', str(self.loja_b.id)]})
        self.assertEqual(lojas_selecionadas(request), [str(self.loja_b.id)])

    def test_sem_filtro_retorna_todas_as_lojas(self):
        nomes = list(filtrar_clientes_solicitacao(self._request({})).values_list('nome', flat=True))
        self.assertCountEqual(nomes, ['Ana Silva', 'Bruno Costa', 'Carla Dias'])

    def test_filtra_por_uma_loja(self):
        request = self._request({'loja': str(self.loja_b.id)})
        nomes = list(filtrar_clientes_solicitacao(request).values_list('nome', flat=True))
        self.assertEqual(nomes, ['Bruno Costa'])

    def test_filtra_por_varias_lojas(self):
        request = self._request({'loja': [str(self.loja_a.id), str(self.loja_c.id)]})
        nomes = list(filtrar_clientes_solicitacao(request).values_list('nome', flat=True))
        self.assertCountEqual(nomes, ['Ana Silva', 'Carla Dias'])

    def test_usuario_sem_permissao_global_ve_apenas_a_loja_da_sessao(self):
        self.user.user_permissions.clear()
        self.user = User.objects.get(pk=self.user.pk)
        request = self._request({'loja': [str(self.loja_b.id), str(self.loja_c.id)]})
        self.assertEqual(list(filtrar_clientes_solicitacao(request)), [])

    def test_combina_filtro_de_loja_com_busca_por_nome(self):
        request = self._request({
            'loja': [str(self.loja_a.id), str(self.loja_b.id)],
            'search': 'Bruno',
        })
        nomes = list(filtrar_clientes_solicitacao(request).values_list('nome', flat=True))
        self.assertEqual(nomes, ['Bruno Costa'])


class ExportacaoClientesExcelTests(FiltroLojaMultiploTests):
    def _planilha(self, params):
        response = ClienteExportExcelView.as_view()(self._request(params))
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        self.assertIn('attachment; filename="clientes_', response['Content-Disposition'])
        return load_workbook(io.BytesIO(response.content)).active

    def test_cabecalho_tem_as_colunas_pedidas(self):
        ws = self._planilha({})
        self.assertEqual(
            [celula.value for celula in ws[1]],
            ['Nome completo', 'CPF', 'Telefone', 'Cidade da solicitação'],
        )

    def test_exporta_dados_formatados_do_cliente(self):
        ws = self._planilha({'loja': str(self.loja_a.id)})
        self.assertEqual(
            list(ws.iter_rows(min_row=2, values_only=True)),
            [('Ana Silva', '123.456.789-00', '(91) 99999-8888', 'Belém')],
        )

    def test_exporta_somente_as_lojas_filtradas(self):
        ws = self._planilha({'loja': [str(self.loja_a.id), str(self.loja_c.id)]})
        nomes = [linha[0] for linha in ws.iter_rows(min_row=2, values_only=True)]
        self.assertCountEqual(nomes, ['Ana Silva', 'Carla Dias'])

    def test_sem_filtro_exporta_toda_a_base_visivel(self):
        ws = self._planilha({})
        self.assertEqual(ws.max_row - 1, 3)

    def test_exportacao_respeita_escopo_de_loja_do_usuario(self):
        self.user.user_permissions.clear()
        self.user.user_permissions.add(Permission.objects.get(codename='view_cliente'))
        self.user = User.objects.get(pk=self.user.pk)
        ws = self._planilha({'loja': str(self.loja_b.id)})
        self.assertEqual(ws.max_row, 1)

    def test_exportacao_ignora_a_paginacao_da_listagem(self):
        for indice in range(15):
            _cliente(f'Cliente {indice:02d}', self.loja_a)
        ws = self._planilha({'loja': str(self.loja_a.id)})
        self.assertEqual(ws.max_row - 1, 16)



class ListagemSolicitacoesRenderTests(TestCase):
    """Garante que a listagem renderiza o filtro múltiplo e o botão de exportar.

    Usa `RequestFactory` em vez do test client porque a instrumentação de
    templates do client quebra nesta combinação Django 4.2 + Python 3.14.
    """

    def setUp(self):
        self.factory = RequestFactory()
        self.loja_a = Loja.objects.create(nome='Loja A')
        self.loja_b = Loja.objects.create(nome='Loja B')
        self.user = User.objects.create_user(
            username='analista', email='analista@teste.com', password='x', loja=self.loja_a,
        )
        self.user.user_permissions.add(
            Permission.objects.get(codename='view_all_analise_credito'),
            Permission.objects.get(codename='view_cliente'),
        )
        self.user = User.objects.get(pk=self.user.pk)
        for indice in range(12):
            _cliente(f'Cliente {indice:02d}', self.loja_a)

    def _html(self, params=None):
        request = self.factory.get('/clientes/', params or {})
        request.user = self.user
        request.session = SessionStore()
        request.session['loja_id'] = self.loja_a.id
        request._messages = FallbackStorage(request)
        response = ClienteListView.as_view()(request)
        self.assertEqual(response.status_code, 200)
        return response.render().content.decode()

    def test_select_de_loja_e_multiplo(self):
        html = self._html()
        self.assertIn('name="loja" id="loja"', html)
        self.assertIn('multiple', html)

    def test_lojas_filtradas_ficam_selecionadas(self):
        html = self._html({'loja': [str(self.loja_a.id), str(self.loja_b.id)]})
        self.assertEqual(html.count('selected>Loja A</option>'), 1)
        self.assertEqual(html.count('selected>Loja B</option>'), 1)

    def test_botao_de_exportar_leva_os_filtros_atuais(self):
        html = self._html({'loja': [str(self.loja_a.id), str(self.loja_b.id)], 'search': 'Ana'})
        self.assertIn('Exportar Excel', html)
        self.assertIn(
            f'/clientes/exportar-excel/?loja={self.loja_a.id}&amp;loja={self.loja_b.id}&amp;search=Ana',
            html,
        )

    def test_paginacao_preserva_as_lojas_selecionadas(self):
        html = self._html({'loja': [str(self.loja_a.id), str(self.loja_b.id)]})
        self.assertIn(
            f'?page=2&amp;loja={self.loja_a.id}&amp;loja={self.loja_b.id}',
            html,
        )
